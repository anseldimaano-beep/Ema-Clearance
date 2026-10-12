import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .models import CAT_LABEL, DEFAULT_WEIGHTS, GRADE_CATS, GradeColumn, GradeIssue, GradeScore, TeacherLoad, infer_cat

SKIP = {"name", "student name", "full name", "last name", "first name", "middle name", "section", "grade level", "grade",
        "remarks", "final", "final grade", "average", "ave", "total", "rank", "sex", "gender", "lrn", "no", "no.",
        "missing", "standing", "standing %", "issues"}
MAX_LABELS = {"max", "highest possible score", "hps", "total score", "perfect score"}
CAT_ROW = {"category", "type", "component"}
ID_RE = re.compile(r"^(student\s*(no\.?|id|number)|id\s*(no\.?|number)?|lrn)$", re.I)
# (name, max score, category) for a brand-new gradebook
DEFAULT_COLS = [("Quiz 1", 20, "quiz"), ("Quiz 2", 20, "quiz"), ("Activity 1", 10, "activity"), ("Activity 2", 10, "activity"),
                ("Project 1", 100, "project"), ("Exam", 50, "exam")]
RED = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="DDE3EE")


def dec(v):
    """'' / None -> None; number -> Decimal; anything else raises ValueError."""
    if v is None or isinstance(v, bool):
        if isinstance(v, bool):
            raise ValueError
        return None
    s = str(v).replace(",", "").strip()
    if not s:
        return None
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError
    if not d.is_finite():
        raise ValueError
    return d.quantize(Decimal("0.01"))


def _norm(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


def cat_key(t):
    """'Quiz' / 'quiz' / 'Recitation / Attendance' / 'attendance' -> category key, else None."""
    s = re.sub(r"[^a-z]", "", str(t or "").lower())
    if not s:
        return None
    for k, lab in GRADE_CATS:
        if s in (k, re.sub(r"[^a-z]", "", lab.lower())):
            return k
    return {"attendance": "recitation", "participation": "recitation", "quizzes": "quiz", "activities": "activity",
            "projects": "project", "exams": "exam"}.get(s)


def parse(fileobj):
    ws = load_workbook(fileobj, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and ID_RE.match(c.strip()) for c in r)), None)
    if hi is None:
        raise ValueError("no header row with a 'Student No' column was found")
    hdr = rows[hi]
    idc = next(j for j, c in enumerate(hdr) if isinstance(c, str) and ID_RE.match(c.strip()))
    acts = [(j, str(h).strip()) for j, h in enumerate(hdr)
            if h is not None and j != idc and str(h).strip() and str(h).strip().lower() not in SKIP]
    if not acts:
        raise ValueError("no activity columns were found")
    maxes, cats, data, texts = {}, {}, {}, 0
    for r in rows[hi + 1:]:
        if not r or all(c is None for c in r):
            continue
        no = _norm(r[idc] if idc < len(r) else None)
        if no.lower() in MAX_LABELS:
            for j, a in acts:
                try:
                    maxes[a] = dec(r[j] if j < len(r) else None)
                except ValueError:
                    pass
            continue
        if no.lower() in CAT_ROW:  # optional row: Quiz / Activity / Project / Exam ... above each activity
            for j, a in acts:
                k = cat_key(r[j] if j < len(r) else None)
                if k:
                    cats[a] = k
            continue
        if not no:
            continue
        rec = {}
        for j, a in acts:
            try:
                rec[a] = dec(r[j] if j < len(r) else None)
            except ValueError:
                rec[a] = None; texts += 1
        data[no] = rec
    return {"acts": [a for _, a in acts], "maxes": maxes, "cats": cats, "data": data, "texts": texts}


def import_into(sheet, parsed):
    loads = {l.student.student_no: l.student for l in
             TeacherLoad.objects.filter(teacher=sheet.teacher, subject=sheet.subject).select_related("student")}
    cols = {}
    for a in parsed["acts"]:
        cat = parsed["cats"].get(a)
        c, _ = GradeColumn.objects.get_or_create(sheet=sheet, name=a,
                                                 defaults={"order": sheet.columns.count(), "category": cat or infer_cat(a)})
        upd = []
        if parsed["maxes"].get(a) is not None:
            c.max_score = parsed["maxes"][a]; upd.append("max_score")
        if cat and c.category != cat:  # only an explicit CATEGORY row overrides what the teacher already set
            c.category = cat; upd.append("category")
        if upd:
            c.save(update_fields=upd)
        cols[a] = c
    matched = 0
    for no, rec in parsed["data"].items():
        st = loads.get(no)
        if not st:
            continue
        matched += 1
        for a, v in rec.items():
            GradeScore.objects.update_or_create(column=cols[a], student=st, defaults={"value": v})
    recompute(sheet)
    return matched, len(parsed["data"]) - matched


def recompute(sheet, student=None):
    cols = list(sheet.columns.all())
    loads = TeacherLoad.objects.filter(teacher=sheet.teacher, subject=sheet.subject).select_related("student")
    old = GradeIssue.objects.filter(sheet=sheet)
    if student:
        loads, old = loads.filter(student=student), old.filter(student=student)
    old.delete()
    sc = {(s.student_id, s.column_id): s.value for s in GradeScore.objects.filter(column__sheet=sheet)}
    new = []
    for l in loads:
        for c in cols:
            v = sc.get((l.student_id, c.id))
            if v is None:
                k, d = "missing", "No score entered"
            elif v < 0:
                k, d = "invalid", "Negative score"
            elif c.max_score is not None and v > c.max_score:
                k, d = "over_max", f"{v:g} is higher than the maximum {c.max_score:g}"
            else:
                continue
            new.append(GradeIssue(sheet=sheet, student=l.student, student_no=l.student.student_no, activity=c.name, kind=k, detail=d))
    GradeIssue.objects.bulk_create(new)


# ---------- grading system ----------

def weights_of(sheet):
    w = {k: float(v) for k, v in DEFAULT_WEIGHTS.items()}
    for k, v in (sheet.weights or {}).items():
        if k in w:
            try:
                w[k] = max(0.0, min(100.0, float(v)))
            except (TypeError, ValueError):
                pass
    return w


def standings(sheet, cols=None, sc=None, ids=None):
    """{student_id: percent or None}. Weighted by category over the scores entered so far; missing work is not counted
    (it is flagged separately), and categories the class does not use are left out and the rest scaled to 100%."""
    cols = list(sheet.columns.all()) if cols is None else cols
    if sc is None:
        sc = {(s.student_id, s.column_id): s.value for s in GradeScore.objects.filter(column__sheet=sheet)}
    if ids is None:
        ids = [l.student_id for l in TeacherLoad.objects.filter(teacher=sheet.teacher, subject=sheet.subject)]
    w = weights_of(sheet)
    out = {}
    for sid in ids:
        num = den = 0.0
        for k, _ in GRADE_CATS:
            if not w[k]:
                continue
            e = p = 0.0
            for c in cols:
                if c.category != k or not c.max_score:
                    continue
                v = sc.get((sid, c.id))
                if v is None or v < 0:
                    continue
                e += float(min(v, c.max_score)); p += float(c.max_score)
            if p:
                num += w[k] * e / p; den += w[k]
        out[sid] = round(100 * num / den, 1) if den else None
    return out


def missing_map(sheet):
    """{student: [activity names still missing, in gradebook order]} for active students."""
    order = {c.name: i for i, c in enumerate(sheet.columns.all())}
    m = {}
    for i in sheet.issues.filter(kind="missing", student__isnull=False, student__is_active=True).select_related("student"):
        m.setdefault(i.student, []).append(i.activity)
    for items in m.values():
        items.sort(key=lambda a: order.get(a, 999))
    return m


def notice_text(subject, teacher, items, note=""):
    """One notification (max 200 chars) listing what the student still has to submit."""
    n = len(items)
    head = f"{subject}: {n} missing requirement{'s' if n != 1 else ''} – "
    tail = f". Please submit to {teacher.display}." + (f" {note}" if note else "")
    room = 200 - len(head) - len(tail)
    for k in range(n, 0, -1):
        body = ", ".join(items[:k]) + (f" (+{n - k} more)" if k < n else "")
        if len(body) <= room:
            return (head + body + tail)[:200]
    return (head + f"{n} item{'s' if n != 1 else ''}" + tail)[:200]


# ---------- Excel ----------

def _title(t, default="Grades"):
    return re.sub(r"[\[\]:*?/\\]", "-", t or default)[:30]


def _book(sheet_title, cols, students, scores=None, stand=None, miss=None):
    """cols: [(name, max, category, column_id)]. With scores, blank cells are shaded red and Missing / Standing % are added."""
    wb = Workbook(); ws = wb.active; ws.title = _title(sheet_title)
    extra = stand is not None
    ws.append(["Student No", "Name"] + [c[0] for c in cols] + (["Missing", "Standing %"] if extra else []))
    ws.append(["CATEGORY", ""] + [CAT_LABEL.get(c[2], "Other") for c in cols])
    ws.append(["MAX", "Highest possible score"] + [c[1] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.fill = HEAD; cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in (2, 3):
        for cell in ws[r]:
            cell.font = Font(italic=True, color="555555")
    for s in students:
        vals = [float(scores[(s.id, c[3])]) if scores and scores.get((s.id, c[3])) is not None else None for c in cols]
        ws.append([s.student_no, s.display] + vals + ([(miss or {}).get(s.id, 0), stand.get(s.id)] if extra else []))
        if extra:
            for j, v in enumerate(vals):
                if v is None:
                    ws.cell(row=ws.max_row, column=3 + j).fill = RED
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 28
    for j in range(len(cols) + 2):
        ws.column_dimensions[ws.cell(row=1, column=3 + j).column_letter].width = 14
    ws.freeze_panes = "C4"
    return wb


def _missing_sheet(wb, sheet, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = "Missing work"
    ws.append(["Student No", "Name", "Requirement", "Category", "Issue", "Detail"])
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.fill = HEAD
    cat = {c.name: c.category for c in sheet.columns.all()}
    order = {c.name: i for i, c in enumerate(sheet.columns.all())}
    iss = sorted(sheet.issues.filter(student__isnull=False).select_related("student"),
                 key=lambda i: (i.student.last_name, i.student.first_name, order.get(i.activity, 999)))
    for i in iss:
        ws.append([i.student_no, i.student.display, i.activity, CAT_LABEL.get(cat.get(i.activity), "Other"),
                   GradeIssue.LABEL.get(i.kind, i.kind).capitalize(), i.detail])
    for col, w in zip("ABCDEF", (16, 28, 22, 22, 14, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws


def _studs(sheet):
    return [l.student for l in TeacherLoad.objects.filter(teacher=sheet.teacher, subject=sheet.subject)
            .select_related("student").order_by("student__last_name", "student__first_name")]


def template_xlsx(subject, students):
    wb = _book(subject, [(n, m, c, None) for n, m, c in DEFAULT_COLS], students)
    b = BytesIO(); wb.save(b); return b.getvalue()


def export_xlsx(sheet):
    studs = _studs(sheet)
    cols = [(c.name, float(c.max_score) if c.max_score is not None else None, c.category, c.id) for c in sheet.columns.all()]
    sc = {(s.student_id, s.column_id): s.value for s in GradeScore.objects.filter(column__sheet=sheet) if s.value is not None}
    miss = {}
    for sid in sheet.issues.filter(kind="missing").values_list("student_id", flat=True):
        miss[sid] = miss.get(sid, 0) + 1
    wb = _book(sheet.subject, cols, studs, sc, standings(sheet, ids=[s.id for s in studs]), miss)
    _missing_sheet(wb, sheet)
    b = BytesIO(); wb.save(b); return b.getvalue()


def missing_xlsx(sheet):
    wb = Workbook()
    _missing_sheet(wb, sheet, first=True)
    b = BytesIO(); wb.save(b); return b.getvalue()
